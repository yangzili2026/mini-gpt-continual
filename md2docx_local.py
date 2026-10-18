# -*- coding: utf-8 -*-
"""Markdown -> docx converter for the two paper drafts (python-docx route)."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_cjk(style, ascii_font, ea_font):
    style.font.name = ascii_font
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), ascii_font)
    rfonts.set(qn('w:hAnsi'), ascii_font)
    rfonts.set(qn('w:eastAsia'), ea_font)

def add_runs(p, text, base_bold=False):
    # split on **bold**
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            r = p.add_run(part); r.bold = base_bold

def flush_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = 'Table Grid'
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            txt = row[j] if j < len(row) else ''
            txt = txt.replace('**', '')
            cell_p = cells[j].paragraphs[0]
            add_runs(cell_p, txt, base_bold=(i == 0))
    doc.add_paragraph()

def convert(md_path, docx_path):
    doc = Document()
    st = doc.styles['Normal']
    st.font.size = Pt(11)
    set_cjk(st, 'Times New Roman', 'SimSun')
    for hname, sz in [('Heading 1', 18), ('Heading 2', 15), ('Heading 3', 13)]:
        hs = doc.styles[hname]
        hs.font.size = Pt(sz)
        hs.font.color.rgb = RGBColor(0, 0, 0)
        set_cjk(hs, 'Times New Roman', 'SimHei')

    lines = open(md_path, encoding='utf-8').read().split('\n')
    table_buf, in_code = [], False
    for line in lines:
        if line.strip().startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.font.name = 'Consolas'
            continue
        s = line.rstrip()
        if s.strip().startswith('|') and s.strip().endswith('|'):
            cells = [c.strip() for c in s.strip().strip('|').split('|')]
            if all(re.fullmatch(r':?-{2,}:?', c or '---') for c in cells):
                continue  # separator row
            table_buf.append(cells)
            continue
        else:
            flush_table(doc, table_buf); table_buf = []
        if not s.strip():
            continue
        if s.strip() == '---':
            continue
        if s.startswith('### '):
            doc.add_heading(s[4:].strip(), level=3)
        elif s.startswith('## '):
            doc.add_heading(s[3:].strip(), level=2)
        elif s.startswith('# '):
            h = doc.add_heading(s[2:].strip(), level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif s.startswith('> '):
            p = doc.add_paragraph(style='Intense Quote' if 'Intense Quote' in [x.name for x in doc.styles] else None)
            add_runs(p, s[2:].strip())
        elif re.match(r'^\s*-\s+', s):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, re.sub(r'^\s*-\s+', '', s))
        else:
            p = doc.add_paragraph()
            add_runs(p, s.strip())
    flush_table(doc, table_buf)
    doc.save(docx_path)
    print('saved:', docx_path)

if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
