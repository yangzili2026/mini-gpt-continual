# -*- coding: utf-8 -*-
"""训练语料素材评分表（全中文素材，面向'喂给自有GPT训练'的视角）→ Word 文档"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 评分维度：中文纯度20% + 知识密度25% + 文本规范性20% + 系统性20% + 训练适配度15%
W = {"中文纯度": 0.20, "知识密度": 0.25, "文本规范性": 0.20, "系统性": 0.20, "训练适配度": 0.15}

MATERIALS = [
    # 名称, 类型, 语言确认依据, 链接, 评分 dict, 点评
    ("《神经网络与深度学习》邱锡鹏（蒲公英书）", "中文原创教材",
     "邱锡鹏中文原创著作，机械工业出版社出版，GitHub 开源中文版全文（nndl 组织）",
     "https://github.com/nndl",
     {"中文纯度": 5.0, "知识密度": 4.5, "文本规范性": 4.8, "系统性": 4.9, "训练适配度": 4.0},
     "中文原创、行文规范、覆盖注意力机制与序列生成模型，是优质的中文深度学习基础语料；缺点是偏理论、GPT 专项内容需到第 8/15 章。"),
    ("《动手学深度学习》李沐等（PyTorch版）", "中文版教材",
     "人民邮电出版社官网标注『语种：中文版』，ISBN 978-7-115-60082-0",
     "https://zh-v2.d2l.ai/",
     {"中文纯度": 4.6, "知识密度": 4.3, "文本规范性": 4.5, "系统性": 4.8, "训练适配度": 4.2},
     "交互式教程、章节完整、术语翻译统一，训练价值高；中英术语混排较多，清洗时需统一处理。"),
    ("《深度学习》Goodfellow 等（花书中文版）", "翻译版名著",
     "人民邮电出版社出版中文版，多所高校教学大纲列为指定教材",
     "https://www.ptpress.com.cn/",
     {"中文纯度": 4.4, "知识密度": 4.8, "文本规范性": 4.6, "系统性": 4.9, "训练适配度": 3.6},
     "经典名著、理论密度极高，适合增强模型数学与机器学习功底；但年代较早、无 GPT 内容，且翻译腔较重，需评估是否合你的训练目标。"),
    ("《从零构建大模型》Raschka（中文版/汉化）", "翻译版专著+代码",
     "MLNLP-World 维护的完整中文翻译仓库，逐章中文笔记与 Jupyter 代码",
     "https://github.com/MLNLP-World/LLMs-from-scratch-CN",
     {"中文纯度": 4.2, "知识密度": 4.6, "文本规范性": 4.3, "系统性": 4.8, "训练适配度": 4.9},
     "与『训练自己的小型GPT』目标完全对齐：从分词、注意力到预训练全流程，且体量适中，是最对口的训练素材之一；为社区翻译版，需核对译文质量。"),
    ("Datawhale《动手学LLM》(llms-from-scratch-cn)", "中文开源教程",
     "Datawhale 中文组织原创中文教程，含 GLM4/Llama3/RWKV6 中文架构笔记",
     "https://github.com/datawhalechina/llms-from-scratch-cn",
     {"中文纯度": 4.5, "知识密度": 4.2, "文本规范性": 4.0, "系统性": 4.5, "训练适配度": 4.8},
     "纯中文社区项目、自带可运行 notebook 讲解，口语化中文多、贴近中文技术表达，做中文指令语料很合适；规范性略逊于正式出版物。"),
    ("Datawhale《Happy-LLM：从零开始构建大模型》", "中文开源电子书",
     "Datawhale 中文原创电子书，章节全部中文撰写",
     "https://github.com/datawhalechina/happy-llm",
     {"中文纯度": 4.7, "知识密度": 4.0, "文本规范性": 4.2, "系统性": 4.6, "训练适配度": 4.7},
     "从 NLP 基础讲到预训练/微调全流程，中文行文流畅、风格统一，适合中量级训练；深度略浅于专著。"),
    ("《这就是ChatGPT》Wolfram（中文版）", "翻译版科普书",
     "人民邮电出版社出版中文版（高校教学大纲参考书目确认）",
     "https://www.ptpress.com.cn/",
     {"中文纯度": 4.5, "知识密度": 3.8, "文本规范性": 4.6, "系统性": 3.9, "训练适配度": 3.8},
     "解释 GPT 原理的绝佳中文科普语料，语言清晰；数学深度有限，适合作为原理性补充语料。"),
    ("《大规模语言模型：从理论到实践》张奇等", "中文原创专著",
     "复旦大学张奇团队中文原创，电子工业出版社 2024 年出版",
     "https://item.jd.com/",
     {"中文纯度": 4.8, "知识密度": 4.7, "文本规范性": 4.7, "系统性": 4.8, "训练适配度": 4.5},
     "中文 LLM 领域少见的系统性原创专著，覆盖训练全流程，语言与知识密度俱佳；体量较大，建议分章喂入。"),
    ("LLM101n 中文翻译版（flypythoncom / wdndev）", "翻译版课程",
     "GitHub 两个活跃中文翻译仓库，课程目录、正文均中文化",
     "https://github.com/wdndev/llm101n-zh",
     {"中文纯度": 4.3, "知识密度": 4.4, "文本规范性": 4.1, "系统性": 4.7, "训练适配度": 4.8},
     "Karpathy 亲自设计的『从零造 Storyteller』课程，从 bigram 到 RLHF 全链路；翻译版质量取决于社区维护，建议抽样校对后使用。"),
    ("nanoGPT-Tutorial-CN（中文教程仓库）", "中文开源教程",
     "GitHub 中文 README，汇集中文博客/视频/笔记链接",
     "https://github.com/paramedick/nanoGPT-Tutorial-CN",
     {"中文纯度": 4.4, "知识密度": 3.6, "文本规范性": 3.8, "系统性": 3.5, "训练适配度": 4.4},
     "极好的『素材索引』而非素材本身：帮你快速定位所有中文 nanoGPT 资料，正文量偏少。"),
    ("知乎专栏《从零构建一个nanoGPT》系列", "中文博客文章",
     "知乎中文专栏，正文中文撰写（2023 年系列学习笔记）",
     "https://zhuanlan.zhihu.com/p/639445749",
     {"中文纯度": 4.5, "知识密度": 3.8, "文本规范性": 3.9, "系统性": 3.7, "训练适配度": 4.2},
     "逐段中文讲解 nanoGPT 源码，第一人称学习笔记风格，口语化表达多，适合增加语料多样性；单篇体量小，需多篇拼合。"),
    ("量子位《李飞飞高徒教你从0到1构建GPT》", "中文媒体文章",
     "量子位中文报道，正文全中文",
     "https://www.qbitai.com/2023/01/41445.html",
     {"中文纯度": 4.6, "知识密度": 3.4, "文本规范性": 4.2, "系统性": 3.2, "训练适配度": 3.7},
     "通俗中文报道，适合提炼背景叙述风格；技术深度有限，宜少量点缀，不宜作主语料。"),
    ("博客园/CSDN《NanoGPT源码详解和中文GPT训练实践》", "中文技术博客",
     "博客园、CSDN 中文技术社区，正文中文并附中文代码注释",
     "https://www.cnblogs.com/justLittleStar/p/18379771",
     {"中文纯度": 4.3, "知识密度": 4.0, "文本规范性": 3.7, "系统性": 3.6, "训练适配度": 4.1},
     "中文注释版源码讲解，工程细节多；博客体例杂、含广告/导航噪音，入库前需清洗。"),
]

def overall(scores):
    return round(sum(scores[k] * W[k] for k in W), 1)

def main():
    out_dir = os.path.dirname(os.path.abspath(__file__))
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)

    title = doc.add_heading("训练语料素材评分表：训练自己的小型 GPT（全中文）", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("用途：将作为语料喂给自有 GPT 训练，评分视角为『语料质量』而非『自学难度』。")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "评分维度与权重：中文纯度 20%（是否全中文、译文质量）＋ 知识密度 25%（单位文本信息量）＋ "
        "文本规范性 20%（书面化程度、噪音多少、是否易清洗）＋ 系统性 20%（体系是否完整、章节是否连贯）＋ "
        "训练适配度 15%（与『从零训练小型 GPT』目标的对口程度）。每项 1–5 分。"
    )

    ranked = sorted(MATERIALS, key=lambda m: -overall(m[4]))
    for rank, (name, mtype, confirm, link, scores, comment) in enumerate(ranked, 1):
        ov = overall(scores)
        doc.add_heading(f"第 {rank} 名｜{name}", level=2)
        meta = doc.add_paragraph()
        meta.add_run(f"类型：{mtype}").bold = True
        doc.add_paragraph(f"语言确认：{confirm}")
        doc.add_paragraph(f"链接：{link}")
        score_line = doc.add_paragraph()
        run = score_line.add_run(
            " ｜ ".join(f"{k} {scores[k]}" for k in W) + f" ｜ ★ 综合 {ov}"
        )
        run.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        doc.add_paragraph(f"点评：{comment}")

    doc.add_heading("使用建议", level=1)
    for tip in [
        "主语料建议以《从零构建大模型》中文版 + Datawhale 两个中文项目为骨架，再用《大规模语言模型：从理论到实践》加深厚度。",
        "博客/媒体类文章务必先清洗网页噪音（导航、广告、评论区）再入库。",
        "翻译版素材建议抽样人工校对 2–3 章，确认术语一致后再整本使用。",
        "注意版权：以上多为受版权保护的出版物，训练使用请确认授权范围，个人学习与内部研究用途风险较低，对外发布模型则需谨慎。",
    ]:
        doc.add_paragraph(tip, style="List Bullet")

    docx_path = os.path.join(out_dir, "小型GPT训练语料素材评分表_全中文.docx")
    doc.save(docx_path)
    print(docx_path)
    for rank, m in enumerate(ranked, 1):
        print(rank, m[0], overall(m[4]))

if __name__ == "__main__":
    main()
